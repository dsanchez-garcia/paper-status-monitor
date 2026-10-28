from __future__ import annotations

import logging
import os
import time
from datetime import datetime
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from paper_status_monitor import storage
from paper_status_monitor.models import COLUMNS, DETECTED_AT_COLUMN, TRACKER_URL_COLUMN
from paper_status_monitor.storage import (
    MAX_STEM_LENGTH,
    SHEET_TITLE,
    WorkbookLockedError,
    WorkbookStore,
    append_record,
    belongs_to,
    find_workbook,
    find_workbook_by_uuid,
    is_named_after,
    last_record,
    merge_workbooks,
    read_records,
    sanitize_filename,
    workbook_name,
)

URL = "https://track.authorhub.elsevier.com/?uuid=aaaaaaaa-bbbb-cccc-dddd-eeeeeeeeeeee"
UUID_A = "aaaaaaaa-bbbb-cccc-dddd-eeeeeeeeeeee"
UUID_B = "11111111-2222-3333-4444-555555555555"
URL_B = f"https://track.authorhub.elsevier.com/?uuid={UUID_B}"


def record(**values: object) -> dict[str, object]:
    base: dict[str, object] = {
        DETECTED_AT_COLUMN: datetime(2026, 1, 2, 9, 30),
        "Title": "A study",
        "Revision": "0",
        "Status": "Under Review",
        "Status code": "3",
        "Reviews completed": "0",
        "Review invitations sent": "2+",
        "Manuscript number": "ABC-D-26-00001",
        TRACKER_URL_COLUMN: URL,
    }
    base.update(values)
    return base


def test_sanitize_filename() -> None:
    assert sanitize_filename('What? A "great": paper/study*  ') == "What A great paperstudy"
    assert sanitize_filename("Ends with dots...") == "Ends with dots"


def test_workbook_name() -> None:
    assert workbook_name("ABC-D-26-00001", "Journal: X", "A study") == (
        "ABC-D-26-00001 - Journal X - A study.xlsx"
    )
    assert workbook_name("ABC-D-26-00001") == "ABC-D-26-00001.xlsx"


def test_workbook_name_is_truncated_on_a_word_boundary() -> None:
    name = workbook_name("ABC-D-26-00001", "Journal", "word " * 100)
    stem = name.removesuffix(".xlsx")
    assert len(stem) <= MAX_STEM_LENGTH
    assert stem.startswith("ABC-D-26-00001 - Journal - word")
    assert stem.endswith("word")


def test_workbook_name_requires_a_key() -> None:
    with pytest.raises(ValueError, match="key"):
        workbook_name("???")


def test_append_creates_a_formatted_workbook(tmp_path: Path) -> None:
    path = tmp_path / "data" / "ABC-D-26-00001 - J - A study.xlsx"
    append_record(path, record())

    workbook = load_workbook(path)
    sheet = workbook[SHEET_TITLE]
    header = [cell.value for cell in sheet[1]]
    assert tuple(header) == COLUMNS
    assert sheet.freeze_panes == "A2"
    values = {label: sheet.cell(2, index + 1).value for index, label in enumerate(header)}
    assert values["Status code"] == 3  # numeric columns are real numbers
    assert values["Reviews completed"] == 0
    assert values["Review invitations sent"] == "2+"
    assert values["Manuscript number"] == "ABC-D-26-00001"
    assert values[DETECTED_AT_COLUMN] == datetime(2026, 1, 2, 9, 30)
    assert sheet.cell(2, header.index(TRACKER_URL_COLUMN) + 1).hyperlink.target == URL
    workbook.close()
    assert not list(path.parent.glob(".*.tmp-*"))  # temporary file cleaned up


def test_read_records_returns_strings(tmp_path: Path) -> None:
    path = tmp_path / "book.xlsx"
    append_record(path, record())
    append_record(path, record(**{"Reviews completed": "1", DETECTED_AT_COLUMN: None}))
    records = read_records(path)
    assert len(records) == 2
    assert records[0]["Status code"] == "3"
    assert records[0][DETECTED_AT_COLUMN] == "2026-01-02 09:30:00"
    assert records[1]["Reviews completed"] == "1"
    assert records[1][DETECTED_AT_COLUMN] == ""
    assert last_record(path) == records[1]


def test_appending_to_a_legacy_workbook_adds_missing_columns(tmp_path: Path) -> None:
    path = tmp_path / "legacy.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Title", "Status", "Manuscript number"])
    sheet.append(["Old title", "Under Review", "ABC-D-26-00001"])
    workbook.save(path)

    append_record(path, record(Title="New title"))

    records = read_records(path)
    assert records[0] == {
        "Title": "Old title",
        "Status": "Under Review",
        "Manuscript number": "ABC-D-26-00001",
        **{label: "" for label in COLUMNS if label not in {"Title", "Status", "Manuscript number"}},
    }
    assert records[1]["Title"] == "New title"
    assert records[1]["Status code"] == "3"
    header = [cell.value for cell in load_workbook(path).worksheets[0][1]]
    assert header[:3] == ["Title", "Status", "Manuscript number"]
    assert set(COLUMNS) <= set(header)


def test_empty_workbook_has_no_last_record(tmp_path: Path) -> None:
    path = tmp_path / "empty.xlsx"
    Workbook().save(path)
    assert read_records(path) == []
    assert last_record(path) is None


def test_find_workbook(tmp_path: Path) -> None:
    assert find_workbook(tmp_path / "missing", "ABC-D-26-00001") is None
    (tmp_path / "ABC-D-26-00001 - Journal - Old title.xlsx").touch()
    (tmp_path / "ABC-D-26-000012 - Journal - Other.xlsx").touch()
    (tmp_path / "~$ABC-D-26-00001 - Journal - Old title.xlsx").touch()
    (tmp_path / "ABC-D-26-00001 - notes.txt").touch()
    found = find_workbook(tmp_path, "ABC-D-26-00001")
    assert found is not None
    assert found.name == "ABC-D-26-00001 - Journal - Old title.xlsx"
    assert find_workbook(tmp_path, "abc-d-26-00001") == found
    assert find_workbook(tmp_path, "XYZ-D-26-00001") is None


def test_find_workbook_prefers_the_most_recent_duplicate(
    tmp_path: Path, caplog: pytest.LogCaptureFixture
) -> None:
    older = tmp_path / "ABC-D-26-00001 - Journal - Title.xlsx"
    newer = tmp_path / "ABC-D-26-00001 - Journal - Title (1).xlsx"
    older.touch()
    newer.touch()
    past = time.time() - 3600
    os.utime(older, (past, past))
    with caplog.at_level(logging.WARNING):
        assert find_workbook(tmp_path, "ABC-D-26-00001") == newer
    assert "Found 2 workbooks" in caplog.text


def test_find_workbook_by_uuid(tmp_path: Path) -> None:
    append_record(tmp_path / "one.xlsx", record())
    append_record(tmp_path / "two.xlsx", record(**{TRACKER_URL_COLUMN: "https://example.org"}))
    (tmp_path / "broken.xlsx").write_text("not a workbook", encoding="utf-8")
    found = find_workbook_by_uuid(tmp_path, "AAAAAAAA-bbbb-cccc-dddd-eeeeeeeeeeee")
    assert found == tmp_path / "one.xlsx"
    assert find_workbook_by_uuid(tmp_path, "11111111-2222-3333-4444-555555555555") is None


def test_locked_workbook_raises_a_clear_error(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    def locked(src: object, dst: object) -> None:
        raise PermissionError("in use")

    monkeypatch.setattr(storage.os, "replace", locked)
    monkeypatch.setattr(storage, "_REPLACE_DELAY", 0)
    path = tmp_path / "book.xlsx"
    with pytest.raises(WorkbookLockedError, match="open in Excel"):
        append_record(path, record())
    assert not path.exists()
    assert not list(tmp_path.glob(".*.tmp-*"))


# ------------------------------------------------------------------ finding histories


def record_b(**values: object) -> dict[str, object]:
    return record(**{"Manuscript number": "XYZ-D-26-00002", TRACKER_URL_COLUMN: URL_B, **values})


def test_belongs_to() -> None:
    row = {"Manuscript number": "ABC-D-26-00001R2", TRACKER_URL_COLUMN: URL}
    assert belongs_to(row, "abc-d-26-00001")
    assert belongs_to(row, "", UUID_A.upper())
    assert not belongs_to(row, "ABC-D-26-0000", UUID_B)
    assert not belongs_to({}, "ABC-D-26-00001", UUID_A)


def test_is_named_after() -> None:
    assert is_named_after(Path("ABC-D-26-00001.xlsx"), "ABC-D-26-00001")
    assert is_named_after(Path("abc-d-26-00001 - J - T.xlsx"), "ABC-D-26-00001")
    assert is_named_after(Path("ABC-D-26-00001 (1).xlsx"), "ABC-D-26-00001")
    assert not is_named_after(Path("ABC-D-26-000012.xlsx"), "ABC-D-26-00001")
    assert not is_named_after(Path("TMY.xlsx"), "")


def test_locate_a_new_submission(tmp_path: Path) -> None:
    location = WorkbookStore(tmp_path / "missing").locate("ABC-D-26-00001", UUID_A, "TMY.xlsx")
    assert location.path == tmp_path / "missing" / "TMY.xlsx"
    assert location.last_record is None


def test_locate_uses_the_workbook_named_after_the_key(tmp_path: Path) -> None:
    old = tmp_path / "ABC-D-26-00001 - J - Old title.xlsx"
    append_record(old, record())
    location = WorkbookStore(tmp_path).locate(
        "ABC-D-26-00001", UUID_A, "ABC-D-26-00001 - J - New title.xlsx"
    )
    assert location.path == old
    assert location.last_record is not None
    assert location.last_record["Status"] == "Under Review"


def test_locate_finds_a_history_by_its_content(tmp_path: Path) -> None:
    renamed = tmp_path / "Whatever the user called it.xlsx"
    append_record(renamed, record())
    append_record(tmp_path / "Unrelated.xlsx", record_b())
    assert WorkbookStore(tmp_path).locate("ABC-D-26-00001", UUID_A, "TMY.xlsx").path == renamed
    # Manuscript number unknown: the UUID is the key.
    assert WorkbookStore(tmp_path).locate(UUID_A, UUID_A, "TMY.xlsx").path == renamed


def test_locate_in_a_shared_workbook(tmp_path: Path) -> None:
    shared = tmp_path / "TMY.xlsx"
    append_record(shared, record(Status="Under Review"))
    append_record(shared, record_b(Status="Review Complete"))
    append_record(shared, record(Status="Required Reviews Complete"))
    store = WorkbookStore(tmp_path)
    first = store.locate("ABC-D-26-00001", UUID_A, "TMY.xlsx")
    second = store.locate("XYZ-D-26-00002", UUID_B, "TMY.xlsx")
    third = store.locate("NEW-D-26-00003", "", "TMY.xlsx")
    assert first.path == second.path == third.path == shared
    assert first.last_record is not None
    assert first.last_record["Status"] == "Required Reviews Complete"
    assert second.last_record is not None
    assert second.last_record["Status"] == "Review Complete"
    assert third.last_record is None  # a new submission that will share the workbook


def test_an_existing_history_wins_over_the_template(tmp_path: Path) -> None:
    append_record(tmp_path / "TMY.xlsx", record_b())  # the target holds another submission
    old = tmp_path / "ABC-D-26-00001 - J - T.xlsx"
    append_record(old, record())
    assert WorkbookStore(tmp_path).locate("ABC-D-26-00001", UUID_A, "TMY.xlsx").path == old


def test_locate_uses_the_hint(tmp_path: Path) -> None:
    hinted = tmp_path / "Hinted.xlsx"
    append_record(hinted, record())
    store = WorkbookStore(tmp_path)
    assert store.locate("ABC-D-26-00001", UUID_A, "TMY.xlsx", hint="Hinted.xlsx").path == hinted
    assert store.locate("ABC-D-26-00001", UUID_A, "TMY.xlsx", hint="Gone.xlsx").path == hinted


def test_rows_without_identifiers_belong_to_the_workbook_named_after_the_key(
    tmp_path: Path,
) -> None:
    path = tmp_path / "ABC-D-26-00001 - J - T.xlsx"
    workbook = Workbook()
    workbook.active.append(["Title", "Status"])
    workbook.active.append(["A study", "Under Review"])
    workbook.save(path)
    location = WorkbookStore(tmp_path).locate("ABC-D-26-00001", UUID_A, "Other.xlsx")
    assert location.last_record == {"Title": "A study", "Status": "Under Review"}
    path.rename(tmp_path / "Legacy.xlsx")
    location = WorkbookStore(tmp_path).locate("ABC-D-26-00001", UUID_A, "Other.xlsx")
    assert location.last_record is None


def test_an_empty_workbook_named_after_the_key_is_reused(tmp_path: Path) -> None:
    empty = tmp_path / "ABC-D-26-00001 - J - T.xlsx"
    Workbook().save(empty)
    location = WorkbookStore(tmp_path).locate("ABC-D-26-00001", UUID_A, "TMY.xlsx")
    assert location.path == empty
    assert location.last_record is None


def test_locate_uuid(tmp_path: Path) -> None:
    shared = tmp_path / "TMY.xlsx"
    append_record(shared, record(Status="Under Review"))
    append_record(shared, record_b())
    append_record(shared, record(**{"Status": "Review Complete", TRACKER_URL_COLUMN: ""}))
    store = WorkbookStore(tmp_path)
    location = store.locate_uuid(UUID_A)
    assert location is not None
    assert location.path == shared
    assert location.key == "ABC-D-26-00001"
    assert location.last_record is not None
    assert location.last_record["Status"] == "Review Complete"  # matched by manuscript number
    assert store.locate_uuid("00000000-0000-0000-0000-000000000000") is None


def test_store_reads_a_workbook_once_until_it_changes(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    path = tmp_path / "book.xlsx"
    append_record(path, record())
    calls: list[Path] = []
    original = storage.read_records

    def counting(target: Path) -> list[dict[str, str]]:
        calls.append(target)
        return original(target)

    monkeypatch.setattr(storage, "read_records", counting)
    store = WorkbookStore(tmp_path)
    store.records(path)
    store.records(path)
    assert len(calls) == 1
    store.append(path, record(Status="Review Complete"))
    assert [row["Status"] for row in store.records(path)] == ["Under Review", "Review Complete"]
    assert len(calls) == 2
    store.forget()
    store.records(path)
    assert len(calls) == 3


def test_unreadable_workbooks_are_skipped(tmp_path: Path) -> None:
    (tmp_path / "broken.xlsx").write_text("not a workbook", encoding="utf-8")
    append_record(tmp_path / "good.xlsx", record())
    store = WorkbookStore(tmp_path)
    assert store.locate("ABC-D-26-00001", UUID_A, "TMY.xlsx").path == tmp_path / "good.xlsx"
    assert store.submission_records(tmp_path / "broken.xlsx", "ABC-D-26-00001") == []


def test_custom_columns_are_appended(tmp_path: Path) -> None:
    path = tmp_path / "book.xlsx"
    append_record(path, {**record(), "Paper": "TMY"})
    header = [cell.value for cell in load_workbook(path)[SHEET_TITLE][1]]
    assert header == [*COLUMNS, "Paper"]
    assert read_records(path)[0]["Paper"] == "TMY"


# ------------------------------------------------------------------ merging


def test_merge_workbooks_in_chronological_order(tmp_path: Path) -> None:
    first = tmp_path / "first.xlsx"
    second = tmp_path / "second.xlsx"
    append_record(first, record(**{DETECTED_AT_COLUMN: datetime(2026, 1, 1), "Status": "a1"}))
    append_record(first, record(**{DETECTED_AT_COLUMN: None, "Status": "a2 (undated)"}))
    append_record(first, record(**{DETECTED_AT_COLUMN: datetime(2026, 3, 1), "Status": "a3"}))
    early = record_b(**{DETECTED_AT_COLUMN: datetime(2026, 2, 1), "Status": "b1"})
    append_record(second, {**early, "Paper": "TMY"})
    append_record(second, record_b(**{DETECTED_AT_COLUMN: datetime(2026, 4, 1), "Status": "b2"}))
    target = tmp_path / "TMY.xlsx"

    assert merge_workbooks(target, [first, second]) == 5

    rows = read_records(target)
    assert [row["Status"] for row in rows] == ["a1", "a2 (undated)", "b1", "a3", "b2"]
    assert rows[2]["Paper"] == "TMY"
    assert first.exists()
    assert second.exists()  # the sources are kept
    sheet = load_workbook(target)[SHEET_TITLE]
    header = [cell.value for cell in sheet[1]]
    assert header == [*COLUMNS, "Paper"]
    assert sheet.cell(2, 1).value == datetime(2026, 1, 1)
    assert sheet.cell(2, header.index(TRACKER_URL_COLUMN) + 1).hyperlink.target == URL
    assert sheet.cell(5, header.index("Status code") + 1).value == 3  # numbers stay numbers


def test_merge_into_an_existing_workbook_keeps_its_other_sheets(tmp_path: Path) -> None:
    target = tmp_path / "TMY.xlsx"
    append_record(target, record(**{DETECTED_AT_COLUMN: datetime(2026, 5, 1), "Status": "new"}))
    workbook = load_workbook(target)
    workbook.create_sheet("Notes")["A1"] = "my notes"
    workbook.save(target)
    source = tmp_path / "old.xlsx"
    append_record(source, record_b(**{DETECTED_AT_COLUMN: datetime(2025, 1, 1), "Status": "old"}))

    assert merge_workbooks(target, [source]) == 2

    assert [row["Status"] for row in read_records(target)] == ["old", "new"]
    assert load_workbook(target)["Notes"]["A1"].value == "my notes"


def test_merge_requires_a_source(tmp_path: Path) -> None:
    with pytest.raises(ValueError, match="source"):
        merge_workbooks(tmp_path / "TMY.xlsx", [])
