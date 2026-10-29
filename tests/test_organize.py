from __future__ import annotations

import os
from dataclasses import replace
from datetime import datetime
from pathlib import Path

import pytest
from openpyxl import load_workbook

from paper_status_monitor import organize
from paper_status_monitor.client import DEMO_UUID, tracker_url
from paper_status_monitor.config import Config, SubmissionConfig
from paper_status_monitor.models import DETECTED_AT_COLUMN, TRACKER_URL_COLUMN
from paper_status_monitor.storage import (
    SHEET_TITLE,
    TIME_SHEET_TITLE,
    WorkbookLockedError,
    append_record,
    read_records,
)

OTHER_UUID = "11111111-2222-3333-4444-555555555555"
THIRD_UUID = "aaaaaaaa-bbbb-cccc-dddd-eeeeeeeeeeee"
NOW = datetime(2026, 10, 28, 9, 0, 0)
FIRST = "DEF-D-26-00001 - Annals of Examples - A study of examples.xlsx"
SECOND = "ABC-D-26-00002 - Journal of Examples - A study of examples.xlsx"


def row(uuid: str, number: str, journal: str, when: datetime, status: str = "Under Review"):
    return {
        DETECTED_AT_COLUMN: when,
        "Title": "A study of examples",
        "Status": status,
        "Journal": journal,
        "Date of submission": "3rd March 2026",
        "Manuscript number": number,
        TRACKER_URL_COLUMN: tracker_url(uuid),
    }


@pytest.fixture
def data_dir(tmp_path: Path) -> Path:
    """Two submissions of one paper, recorded with the default names."""
    folder = tmp_path / "data"
    first = folder / FIRST
    append_record(
        first, row(DEMO_UUID, "DEF-D-26-00001", "Annals of Examples", datetime(2026, 1, 5))
    )
    append_record(
        first,
        row(
            DEMO_UUID,
            "DEF-D-26-00001",
            "Annals of Examples",
            datetime(2026, 3, 1),
            "Review Complete",
        ),
    )
    second = folder / SECOND
    append_record(
        second,
        row(OTHER_UUID, "ABC-D-26-00002", "Journal of Examples", datetime(2026, 4, 2)),
    )
    return folder


def make_config(data_dir: Path, template: str) -> Config:
    return Config(
        data_dir=data_dir,
        workbook_name=template,
        submissions=(
            SubmissionConfig(DEMO_UUID, label="Annals", fields={"paper": "Tau"}),
            SubmissionConfig(OTHER_UUID, label="Journal", fields={"paper": "Tau"}, enabled=False),
        ),
    )


def names(folder: Path) -> list[str]:
    return sorted(path.name for path in folder.glob("*.xlsx"))


def test_nothing_to_do_when_the_names_follow_the_template(data_dir: Path) -> None:
    plan = organize.plan_organize(
        make_config(data_dir, "{manuscript_number} - {journal} - {title}")
    )
    assert plan.actions == []
    assert plan.skipped == []
    assert [name.current.name for name in plan.names if name.current] == [FIRST, SECOND]
    assert [name.needs_change for name in plan.names] == [False, False]


def test_rename_to_follow_the_template(data_dir: Path) -> None:
    config = make_config(data_dir, "{paper} - {journal_acronym} - {manuscript_number}")
    plan = organize.plan_organize(config)
    assert [action.kind for action in plan.actions] == ["rename", "rename"]
    assert plan.actions[0].describe() == (
        f"Rename  {FIRST}\n    ->  Tau - DEF - DEF-D-26-00001.xlsx"
    )
    assert plan.actions[0].submissions == ("Annals",)

    outcome = organize.apply_organize(plan, config, now=lambda: NOW)

    assert outcome.errors == []
    assert outcome.backup_dir is None  # renaming needs no copies
    assert len(outcome.done) == 2
    assert names(data_dir) == ["Tau - ABC - ABC-D-26-00002.xlsx", "Tau - DEF - DEF-D-26-00001.xlsx"]
    assert organize.plan_organize(config).actions == []


def test_merge_the_submissions_of_a_paper(data_dir: Path) -> None:
    config = make_config(data_dir, "{paper}")
    plan = organize.plan_organize(config)
    [action] = plan.actions
    assert action.kind == "merge"
    assert action.target == data_dir / "Tau.xlsx"
    assert action.sources == (data_dir / FIRST, data_dir / SECOND)
    assert action.describe().splitlines()[0] == "Merge into  Tau.xlsx"

    outcome = organize.apply_organize(plan, config, now=lambda: NOW)

    assert outcome.errors == []
    assert names(data_dir) == ["Tau.xlsx"]
    rows = read_records(data_dir / "Tau.xlsx")
    assert [(r["Manuscript number"], r["Status"]) for r in rows] == [
        ("DEF-D-26-00001", "Under Review"),
        ("DEF-D-26-00001", "Review Complete"),
        ("ABC-D-26-00002", "Under Review"),
    ]
    assert outcome.backup_dir == data_dir / "backups" / "organize-2026-10-28_09-00-00"
    assert sorted(path.name for path in outcome.backup_dir.iterdir()) == [SECOND, FIRST]
    assert organize.plan_organize(config).actions == []


def test_merging_adds_the_time_in_status_sheet(data_dir: Path) -> None:
    config = replace(make_config(data_dir, "{paper}"), time_in_status=True)
    outcome = organize.apply_organize(organize.plan_organize(config), config, now=lambda: NOW)
    assert outcome.errors == []
    workbook = load_workbook(data_dir / "Tau.xlsx")
    assert workbook.sheetnames == [SHEET_TITLE, TIME_SHEET_TITLE]
    sheet = workbook[TIME_SHEET_TITLE]
    assert [sheet.cell(row, 1).value for row in (5, 6, 7)] == [
        "DEF-D-26-00001",
        "DEF-D-26-00001",
        "ABC-D-26-00002",
    ]


def test_merge_into_an_existing_shared_workbook(data_dir: Path) -> None:
    shared = data_dir / "Tau.xlsx"
    append_record(shared, row(THIRD_UUID, "XYZ-D-26-00003", "Journal X", datetime(2026, 9, 1)))
    base = make_config(data_dir, "{paper}")
    third = SubmissionConfig(THIRD_UUID, fields={"paper": "Tau"})
    config = replace(base, submissions=(*base.submissions, third))
    plan = organize.plan_organize(config)
    [action] = plan.actions
    assert action.kind == "merge"
    assert action.sources == (data_dir / FIRST, data_dir / SECOND)
    assert "(existing workbook)" in action.describe()

    outcome = organize.apply_organize(plan, config, now=lambda: NOW)

    assert outcome.errors == []
    assert [r["Manuscript number"] for r in read_records(shared)] == [
        "DEF-D-26-00001",
        "DEF-D-26-00001",
        "ABC-D-26-00002",
        "XYZ-D-26-00003",
    ]
    assert outcome.backup_dir is not None
    assert (outcome.backup_dir / "Tau.xlsx").is_file()


def test_workbooks_with_other_submissions_are_left_alone(data_dir: Path) -> None:
    append_record(
        data_dir / FIRST, row(THIRD_UUID, "XYZ-D-26-00003", "Journal X", datetime(2026, 5, 1))
    )
    plan = organize.plan_organize(make_config(data_dir, "{paper}"))
    [skipped] = plan.skipped
    assert skipped.startswith(FIRST)
    assert "not in the configuration" in skipped
    [action] = plan.actions
    assert (action.kind, action.sources) == ("rename", (data_dir / SECOND,))


def merged(data_dir: Path) -> Config:
    """Merge the two submissions of the fixture into Tau.xlsx; return the configuration."""
    config = make_config(data_dir, "{paper}")
    organize.apply_organize(organize.plan_organize(config), config, now=lambda: NOW)
    assert names(data_dir) == ["Tau.xlsx"]
    return config


def manuscripts(path: Path) -> list[str]:
    return [record["Manuscript number"] for record in read_records(path)]


def test_a_shared_workbook_is_split(data_dir: Path) -> None:
    config = merged(data_dir)
    split = replace(
        config,
        submissions=(
            SubmissionConfig(DEMO_UUID, label="Annals", fields={"paper": "Tau"}),
            SubmissionConfig(OTHER_UUID, label="Journal", fields={"paper": "Other"}),
        ),
        time_in_status=True,
    )
    plan = organize.plan_organize(split)
    [action] = plan.actions
    assert action.kind == "split"
    assert action.describe() == (
        "Split   Tau.xlsx\n    ->  Tau.xlsx (kept here): Annals\n    ->  Other.xlsx: Journal"
    )

    outcome = organize.apply_organize(plan, split, now=lambda: NOW)

    assert outcome.errors == []
    assert outcome.done == ["Split Tau.xlsx: kept 2 rows, moved the others to Other.xlsx (1 row)"]
    assert names(data_dir) == ["Other.xlsx", "Tau.xlsx"]
    assert manuscripts(data_dir / "Tau.xlsx") == ["DEF-D-26-00001", "DEF-D-26-00001"]
    assert manuscripts(data_dir / "Other.xlsx") == ["ABC-D-26-00002"]
    assert load_workbook(data_dir / "Other.xlsx").sheetnames == [SHEET_TITLE, TIME_SHEET_TITLE]
    assert outcome.backup_dir is not None
    assert len(read_records(outcome.backup_dir / "Tau.xlsx")) == 3
    assert organize.plan_organize(split).actions == []


def test_a_workbook_is_split_into_new_workbooks(data_dir: Path) -> None:
    merged(data_dir)
    by_journal = make_config(data_dir, "{paper} - {journal_acronym}")
    plan = organize.plan_organize(by_journal)
    assert [action.kind for action in plan.actions] == ["split"]

    outcome = organize.apply_organize(plan, by_journal, now=lambda: NOW)

    assert outcome.errors == []
    assert outcome.done == ["Split Tau.xlsx into Tau - DEF.xlsx (2 rows), Tau - ABC.xlsx (1 row)"]
    assert names(data_dir) == ["Tau - ABC.xlsx", "Tau - DEF.xlsx"]
    assert manuscripts(data_dir / "Tau - ABC.xlsx") == ["ABC-D-26-00002"]


def test_a_split_merges_the_rows_into_an_existing_workbook(data_dir: Path) -> None:
    config = merged(data_dir)
    other = data_dir / "Other.xlsx"
    append_record(other, row(THIRD_UUID, "XYZ-D-26-00003", "Journal X", datetime(2026, 5, 1)))
    split = replace(
        config,
        submissions=(
            SubmissionConfig(DEMO_UUID, label="Annals", fields={"paper": "Tau"}),
            SubmissionConfig(OTHER_UUID, label="Journal", fields={"paper": "Other"}),
            SubmissionConfig(THIRD_UUID, label="X", fields={"paper": "Other"}),
        ),
    )
    plan = organize.plan_organize(split)
    [action] = plan.actions
    assert action.describe().splitlines()[-1] == "    ->  Other.xlsx (existing workbook): Journal"

    outcome = organize.apply_organize(plan, split, now=lambda: NOW)

    assert outcome.errors == []
    assert manuscripts(other) == ["ABC-D-26-00002", "XYZ-D-26-00003"]
    assert manuscripts(data_dir / "Tau.xlsx") == ["DEF-D-26-00001", "DEF-D-26-00001"]
    assert outcome.backup_dir is not None
    assert manuscripts(outcome.backup_dir / "Other.xlsx") == ["XYZ-D-26-00003"]


def test_workbooks_with_unknown_rows_are_not_split(data_dir: Path) -> None:
    config = merged(data_dir)
    shared = data_dir / "Tau.xlsx"
    append_record(shared, row(THIRD_UUID, "XYZ-D-26-00003", "Journal X", datetime(2026, 5, 1)))
    split = replace(
        config,
        submissions=(
            SubmissionConfig(DEMO_UUID, fields={"paper": "Tau"}),
            SubmissionConfig(OTHER_UUID, fields={"paper": "Other"}),
        ),
    )
    plan = organize.plan_organize(split)
    assert plan.actions == []
    [skipped] = plan.skipped
    assert "1 of its rows belongs to none of them" in skipped


def test_a_workbook_in_use_is_not_split(data_dir: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    config = merged(data_dir)
    split = replace(
        config,
        submissions=(
            SubmissionConfig(DEMO_UUID, fields={"paper": "Tau"}),
            SubmissionConfig(OTHER_UUID, fields={"paper": "Other"}),
        ),
    )

    def in_use(paths: object) -> None:
        raise WorkbookLockedError("Cannot write Tau.xlsx: the file is in use.")

    monkeypatch.setattr(organize.storage, "ensure_writable", in_use)
    outcome = organize.apply_organize(organize.plan_organize(split), split, now=lambda: NOW)
    assert outcome.errors == ["Split Tau.xlsx: Cannot write Tau.xlsx: the file is in use."]
    assert names(data_dir) == ["Tau.xlsx"]
    assert len(read_records(data_dir / "Tau.xlsx")) == 3


def test_chained_renames_are_done_in_two_steps(tmp_path: Path) -> None:
    append_record(tmp_path / "X.xlsx", row(DEMO_UUID, "DEF-D-26-00001", "A", datetime(2026, 1, 1)))
    append_record(tmp_path / "Y.xlsx", row(OTHER_UUID, "ABC-D-26-00002", "J", datetime(2026, 1, 2)))
    config = Config(
        data_dir=tmp_path,
        workbook_name="{paper}",
        submissions=(
            SubmissionConfig(DEMO_UUID, fields={"paper": "Y"}),
            SubmissionConfig(OTHER_UUID, fields={"paper": "Z"}),
        ),
    )
    plan = organize.plan_organize(config)
    assert [action.kind for action in plan.actions] == ["rename"]
    assert "run 'organize' again" in plan.skipped[0]
    organize.apply_organize(plan, config)
    organize.apply_organize(organize.plan_organize(config), config)
    assert names(tmp_path) == ["Y.xlsx", "Z.xlsx"]
    assert read_records(tmp_path / "Y.xlsx")[0]["Manuscript number"] == "DEF-D-26-00001"


def test_failed_actions_are_reported_and_the_others_still_run(
    data_dir: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    config = make_config(data_dir, "{paper} - {journal_acronym}")
    real_replace = os.replace
    calls: list[object] = []

    def flaky(source: object, target: object) -> None:
        calls.append(target)
        if len(calls) == 1:
            raise PermissionError("the file is in use")
        real_replace(source, target)

    monkeypatch.setattr(organize.os, "replace", flaky)
    outcome = organize.apply_organize(organize.plan_organize(config), config)
    assert len(outcome.errors) == 1
    assert "Close the workbooks in Excel" in outcome.errors[0]
    assert len(outcome.done) == 1


def test_submissions_not_recorded_yet(tmp_path: Path) -> None:
    config = Config(data_dir=tmp_path / "empty", submissions=(SubmissionConfig(DEMO_UUID),))
    plan = organize.plan_organize(config)
    assert plan.actions == []
    [name] = plan.names
    assert name.current is None
    assert not name.needs_change
