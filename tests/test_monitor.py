from __future__ import annotations

from dataclasses import replace
from pathlib import Path

import pytest
from openpyxl import load_workbook

from paper_status_monitor import storage
from paper_status_monitor.client import (
    DEMO_UUID,
    SubmissionNotFoundError,
    TrackerUnavailableError,
)
from paper_status_monitor.config import Config, SubmissionConfig
from paper_status_monitor.models import DETECTED_AT_COLUMN, TRACKER_URL_COLUMN, UNAVAILABLE_STATUS
from paper_status_monitor.monitor import Monitor
from paper_status_monitor.state import SEEN_FILE_NAME, SeenSnapshots
from paper_status_monitor.storage import read_records

OTHER_UUID = "11111111-2222-3333-4444-555555555555"
MAY_18 = 1652887284  # "Last review activity" one day after the demo record


def workbooks(config: Config) -> list[Path]:
    return sorted(config.data_dir.glob("*.xlsx"))


def test_time_in_status_sheet_is_added_and_kept_up_to_date(
    config, fake_client, make_payload, clock
) -> None:
    fake_client.queue(DEMO_UUID, make_payload(), make_payload(), make_payload(Status=4))
    Monitor(config, fake_client, now=clock).run()
    [path] = workbooks(config)
    assert load_workbook(path).sheetnames == [storage.SHEET_TITLE]  # the setting is off

    enabled = replace(config, time_in_status=True)
    clock.advance(days=1)
    Monitor(enabled, fake_client, now=clock).run()  # nothing changed: the sheet is added
    assert load_workbook(path).sheetnames == [storage.SHEET_TITLE, storage.TIME_SHEET_TITLE]

    clock.advance(days=2)
    [result] = Monitor(enabled, fake_client, now=clock).run().results  # a change: rewritten
    assert result.recorded
    sheet = load_workbook(path)[storage.TIME_SHEET_TITLE]
    assert [sheet["D5"].value, sheet["G5"].value, sheet["D6"].value] == [
        "Under Review",
        3,
        "Required Reviews Complete",
    ]


def test_time_in_status_problems_do_not_stop_a_check(
    config, fake_client, make_payload, clock, monkeypatch: pytest.MonkeyPatch
) -> None:
    fake_client.queue(DEMO_UUID, make_payload())
    Monitor(config, fake_client, now=clock).run()

    def locked(path: Path) -> bool:
        raise storage.WorkbookLockedError("the file is in use")

    monkeypatch.setattr(storage, "write_time_sheet", locked)
    [result] = Monitor(replace(config, time_in_status=True), fake_client, now=clock).run().results
    assert result.ok
    assert load_workbook(workbooks(config)[0]).sheetnames == [storage.SHEET_TITLE]


def test_first_run_creates_one_workbook(config, fake_client, make_payload, clock) -> None:
    fake_client.queue(DEMO_UUID, make_payload())
    summary = Monitor(config, fake_client, now=clock).run()

    [result] = summary.results
    assert result.ok
    assert result.is_new
    assert result.changes == []
    assert summary.changed == [result]
    [path] = workbooks(config)
    assert path == result.workbook
    assert path.name.startswith(
        "SUBTRAK-D-22-00128 - The International Journal of Submission Trackers - Influence"
    )
    [row] = read_records(path)
    assert row[DETECTED_AT_COLUMN] == "2026-01-02 09:30:00"
    assert row["Status"] == "Under Review"
    assert row[TRACKER_URL_COLUMN].endswith(DEMO_UUID)


def test_unchanged_status_adds_no_row(config, fake_client, make_payload, clock) -> None:
    fake_client.queue(DEMO_UUID, make_payload(), make_payload())
    monitor = Monitor(config, fake_client, now=clock)
    monitor.run()
    summary = monitor.run()

    [result] = summary.results
    assert result.ok
    assert not result.is_new
    assert result.changes == []
    assert summary.changed == []
    assert len(read_records(workbooks(config)[0])) == 1


def test_changes_are_detected_and_appended(config, fake_client, make_payload, clock) -> None:
    fake_client.queue(
        DEMO_UUID,
        make_payload(),
        make_payload(Status=4, LastUpdated=1652887284, ReviewsCompleted_summary="2"),
    )
    monitor = Monitor(config, fake_client, now=clock)
    monitor.run()
    clock.advance(days=1)
    summary = monitor.run()

    [result] = summary.results
    assert [str(change) for change in result.changes] == [
        "Status: 'Under Review' → 'Required Reviews Complete'",
        "Status code: '3' → '4'",
        "Last review activity: '17th May 2022' → '18th May 2022'",
        "Reviews completed: '1' → '2'",
    ]
    rows = read_records(workbooks(config)[0])
    assert [row["Status"] for row in rows] == ["Under Review", "Required Reviews Complete"]
    assert rows[1][DETECTED_AT_COLUMN] == "2026-01-03 09:30:00"


def test_revisions_share_the_same_workbook(config, fake_client, make_payload, clock) -> None:
    fake_client.queue(
        DEMO_UUID,
        make_payload(),
        make_payload(LatestRevisionNumber=1, Status=3, ReviewsCompleted_summary="0"),
    )
    monitor = Monitor(config, fake_client, now=clock)
    monitor.run()
    summary = monitor.run()

    [result] = summary.results
    assert not result.is_new
    fields = [change.field for change in result.changes]
    assert fields == ["Revision", "Reviews completed", "Manuscript number"]
    [path] = workbooks(config)
    rows = read_records(path)
    assert [row["Manuscript number"] for row in rows] == [
        "SUBTRAK-D-22-00128",
        "SUBTRAK-D-22-00128R1",
    ]


def test_ignored_fields_are_recorded_but_not_reported(
    config, fake_client, make_payload, clock
) -> None:
    config = replace(config, ignore_fields=("Status code",))
    fake_client.queue(DEMO_UUID, make_payload(Status=8), make_payload(Status=9))
    monitor = Monitor(config, fake_client, now=clock)
    monitor.run()
    summary = monitor.run()

    assert summary.results[0].changes == []
    assert summary.changed == []
    rows = read_records(workbooks(config)[0])
    assert [row["Status code"] for row in rows] == ["8", "9"]


def test_unknown_uuid_is_an_error(config, fake_client) -> None:
    fake_client.queue(DEMO_UUID, SubmissionNotFoundError("HTTP 404"))
    summary = Monitor(config, fake_client).run()

    [result] = summary.results
    assert not result.ok
    assert "404" in (result.error or "")
    assert summary.failed == [result]
    assert workbooks(config) == []


def test_known_submission_disappearing_is_recorded_once(
    config, fake_client, make_payload, clock
) -> None:
    fake_client.queue(
        DEMO_UUID,
        make_payload(Status=29),
        SubmissionNotFoundError("HTTP 404"),
        SubmissionNotFoundError("HTTP 404"),
    )
    monitor = Monitor(config, fake_client, now=clock)
    monitor.run()
    second = monitor.run()
    third = monitor.run()

    [result] = second.results
    assert result.ok
    assert [str(change) for change in result.changes] == [
        f"Status: 'Review Complete' → '{UNAVAILABLE_STATUS}'",
        "Status code: '29' → ''",
    ]
    assert third.results[0].ok
    assert third.results[0].changes == []
    rows = read_records(workbooks(config)[0])
    assert [row["Status"] for row in rows] == ["Review Complete", UNAVAILABLE_STATUS]


def test_network_errors_do_not_write_anything(config, fake_client) -> None:
    fake_client.queue(DEMO_UUID, TrackerUnavailableError("offline"))
    summary = Monitor(config, fake_client).run()
    assert summary.failed
    assert summary.results[0].error == "offline"
    assert workbooks(config) == []


def test_one_failure_does_not_stop_the_others(config, fake_client, make_payload) -> None:
    config = replace(
        config,
        submissions=(SubmissionConfig(OTHER_UUID), *config.submissions),
    )
    fake_client.queue(OTHER_UUID, TrackerUnavailableError("offline"))
    fake_client.queue(DEMO_UUID, make_payload())
    summary = Monitor(config, fake_client).run()
    assert [result.ok for result in summary.results] == [False, True]
    assert len(workbooks(config)) == 1


def test_unexpected_errors_are_captured(config, fake_client, make_payload, monkeypatch) -> None:
    def explode(*args: object, **kwargs: object) -> None:
        raise RuntimeError("disk on fire")

    monkeypatch.setattr("paper_status_monitor.storage.append_record", explode)
    fake_client.queue(DEMO_UUID, make_payload())
    summary = Monitor(config, fake_client).run()
    assert summary.results[0].error == "RuntimeError: disk on fire"


def test_dry_run_writes_nothing(config, fake_client, make_payload) -> None:
    fake_client.queue(DEMO_UUID, make_payload())
    summary = Monitor(config, fake_client).run(dry_run=True)
    assert summary.results[0].is_new
    assert not config.data_dir.exists()


def test_disabled_submissions_are_skipped(config, fake_client) -> None:
    config = replace(config, submissions=(SubmissionConfig(DEMO_UUID, enabled=False),))
    summary = Monitor(config, fake_client).run()
    assert summary.results == []
    assert fake_client.calls == []


def test_wait_until_ready_creates_the_data_folder(config, fake_client) -> None:
    config.data_dir.parent.mkdir(parents=True, exist_ok=True)
    problems = Monitor(config, fake_client).wait_until_ready(timeout=0)
    assert problems == []
    assert config.data_dir.is_dir()


def test_wait_until_ready_reports_problems_after_the_timeout(tmp_path, fake_client) -> None:
    config = Config(data_dir=tmp_path / "not" / "mounted" / "data")
    fake_client.reachable = False
    sleeps: list[float] = []
    now = iter([0.0, 0.0, 6.0, 12.0])

    problems = Monitor(config, fake_client).wait_until_ready(
        timeout=10, interval=6, sleep=sleeps.append, clock=lambda: next(now)
    )

    assert sleeps == [6, 6]
    assert len(problems) == 2
    assert "data folder" in problems[0]
    assert "tracker" in problems[1]


@pytest.mark.parametrize("label", ["", "My paper"])
def test_result_name_falls_back_to_label_or_uuid(config, fake_client, label) -> None:
    config = replace(config, submissions=(SubmissionConfig(DEMO_UUID, label=label),))
    fake_client.queue(DEMO_UUID, TrackerUnavailableError("offline"))
    [result] = Monitor(config, fake_client).run().results
    assert result.name == (label or DEMO_UUID)


# ------------------------------------------------------------------ record vs notify


def test_changes_can_be_recorded_without_being_notified(
    config, fake_client, make_payload, clock
) -> None:
    """Only 'Last review activity' changed: keep it in the workbook, but do not report it."""
    config = replace(config, notify_ignore=("Last review activity",))
    fake_client.queue(DEMO_UUID, make_payload(), make_payload(LastUpdated=MAY_18))
    monitor = Monitor(config, fake_client, now=clock)
    monitor.run()
    clock.advance(days=1)
    summary = monitor.run()

    [result] = summary.results
    assert result.changes == []
    assert [str(change) for change in result.recorded_changes] == [
        "Last review activity: '17th May 2022' → '18th May 2022'"
    ]
    assert summary.changed == []
    assert summary.recorded == [result]
    rows = read_records(workbooks(config)[0])
    assert [row["Last review activity"] for row in rows] == ["17th May 2022", "18th May 2022"]


def test_only_status_changes_are_notified(config, fake_client, make_payload, clock) -> None:
    config = replace(config, notify_fields=("Status",))
    fake_client.queue(
        DEMO_UUID,
        make_payload(),
        make_payload(LastUpdated=MAY_18, ReviewsCompleted_summary="2"),
        make_payload(Status=4, LastUpdated=MAY_18, ReviewsCompleted_summary="2"),
    )
    monitor = Monitor(config, fake_client, now=clock)
    monitor.run()
    second = monitor.run()
    third = monitor.run()

    assert second.changed == []
    assert [c.field for c in second.results[0].recorded_changes] == [
        "Last review activity",
        "Reviews completed",
    ]
    assert [str(change) for change in third.results[0].changes] == [
        "Status: 'Under Review' → 'Required Reviews Complete'"
    ]
    assert [c.field for c in third.results[0].unreported_changes] == ["Status code"]
    assert len(read_records(workbooks(config)[0])) == 3


def test_each_submission_can_have_its_own_settings(
    config, fake_client, make_payload, clock
) -> None:
    """The general settings notify everything; one submission only reports its status."""
    own = SubmissionConfig(
        OTHER_UUID, label="Own", overrides={"notify_fields": ("Status",), "record_fields": ()}
    )
    config = replace(config, submissions=(*config.submissions, own))
    fake_client.queue(DEMO_UUID, make_payload(), make_payload(LastUpdated=MAY_18))
    other = {"PubdNumber": "ABC-D-26-00001"}
    fake_client.queue(OTHER_UUID, make_payload(**other), make_payload(**other, LastUpdated=MAY_18))
    monitor = Monitor(config, fake_client, now=clock)
    first = monitor.run()
    general, special = monitor.run().results

    assert [result.is_new for result in first.results] == [True, True]
    assert [c.field for c in general.changes] == ["Last review activity"]
    assert [c.field for c in general.recorded_changes] == ["Last review activity"]
    assert special.changes == []  # only changes of status are reported for it
    assert special.recorded_changes == []  # and no change adds a row
    assert len(workbooks(config)) == 2


def test_changes_can_be_notified_without_being_recorded(
    config, fake_client, make_payload, clock
) -> None:
    config = replace(config, record_ignore=("Reviews completed",))
    fake_client.queue(
        DEMO_UUID,
        make_payload(),
        make_payload(ReviewsCompleted_summary="2"),
        make_payload(ReviewsCompleted_summary="2"),
        make_payload(Status=4, ReviewsCompleted_summary="2"),
    )
    monitor = Monitor(config, fake_client, now=clock)
    monitor.run()
    second = monitor.run()
    third = monitor.run()
    fourth = monitor.run()

    assert [str(change) for change in second.results[0].changes] == ["Reviews completed: '1' → '2'"]
    assert second.results[0].recorded_changes == []
    assert third.changed == []  # reported once, not at every run
    assert [c.field for c in fourth.results[0].changes] == ["Status", "Status code"]
    rows = read_records(workbooks(config)[0])
    assert [row["Reviews completed"] for row in rows] == ["1", "2"]  # kept with the next row


def test_only_some_columns_add_rows(config, fake_client, make_payload, clock) -> None:
    config = replace(config, record_fields=("Status",))
    fake_client.queue(
        DEMO_UUID,
        make_payload(),
        make_payload(LastUpdated=MAY_18),
        make_payload(LastUpdated=MAY_18),
    )
    monitor = Monitor(config, fake_client, now=clock)
    monitor.run()
    second = monitor.run()
    third = monitor.run()

    assert [c.field for c in second.results[0].changes] == ["Last review activity"]
    assert not second.results[0].recorded
    assert third.changed == []
    assert len(read_records(workbooks(config)[0])) == 1


def test_last_state_seen_is_kept_in_the_data_folder(
    config, fake_client, make_payload, clock
) -> None:
    fake_client.queue(DEMO_UUID, make_payload(), make_payload(Status=4))
    monitor = Monitor(config, fake_client, now=clock)
    monitor.run()
    seen_file = config.data_dir / SEEN_FILE_NAME
    seen = SeenSnapshots.load(seen_file)
    assert (seen.values(DEMO_UUID) or {})["Status"] == "Under Review"
    assert seen.workbook(DEMO_UUID) == workbooks(config)[0].name

    monitor.run(dry_run=True)
    assert (SeenSnapshots.load(seen_file).values(DEMO_UUID) or {})["Status"] == "Under Review"


def test_check_outside_a_run_saves_the_state(config, fake_client, make_payload, clock) -> None:
    fake_client.queue(DEMO_UUID, make_payload())
    result = Monitor(config, fake_client, now=clock).check(config.submissions[0])
    assert result.is_new
    assert (config.data_dir / SEEN_FILE_NAME).is_file()


def test_state_file_errors_do_not_stop_a_run(
    config, fake_client, make_payload, monkeypatch
) -> None:
    def fail(self: SeenSnapshots) -> bool:
        raise OSError("disk full")

    monkeypatch.setattr(SeenSnapshots, "save", fail)
    fake_client.queue(DEMO_UUID, make_payload())
    summary = Monitor(config, fake_client).run()
    assert summary.results[0].ok
    assert summary.results[0].is_new


# ------------------------------------------------------------------ workbook names


def paper_config(config: Config, template: str = "{paper}") -> Config:
    return replace(
        config,
        workbook_name=template,
        submissions=(
            SubmissionConfig(DEMO_UUID, label="First journal", fields={"paper": "Tau"}),
            SubmissionConfig(OTHER_UUID, label="Second journal", fields={"paper": "Tau"}),
        ),
    )


def other_payload(make_payload, **changes: object) -> dict:
    return make_payload(
        PubdNumber="OTHER-D-26-00042", JournalName="Journal of Second Chances", **changes
    )


def test_the_template_names_new_workbooks(config, fake_client, make_payload, clock) -> None:
    config = replace(
        paper_config(config, "{paper} - {journal_acronym} - {manuscript_number}"),
        submissions=(SubmissionConfig(DEMO_UUID, fields={"paper": "Tau"}),),
    )
    fake_client.queue(DEMO_UUID, make_payload())
    [result] = Monitor(config, fake_client, now=clock).run().results
    assert result.workbook == config.data_dir / "Tau - SUBTRAK - SUBTRAK-D-22-00128.xlsx"
    [row] = read_records(result.workbook)
    assert row["Paper"] == "Tau"  # custom fields are written to extra columns


def test_submissions_of_a_paper_can_share_a_workbook(
    config, fake_client, make_payload, clock
) -> None:
    config = paper_config(config)
    fake_client.queue(DEMO_UUID, make_payload(Status=29), make_payload(Status=29))
    fake_client.queue(
        OTHER_UUID, other_payload(make_payload), other_payload(make_payload, Status=4)
    )
    monitor = Monitor(config, fake_client, now=clock)
    first = monitor.run()
    clock.advance(days=1)
    second = monitor.run()

    assert [result.is_new for result in first.results] == [True, True]
    [path] = workbooks(config)
    assert path.name == "Tau.xlsx"
    assert [result.workbook for result in second.results] == [path, path]
    assert second.results[0].changes == []
    assert [c.field for c in second.results[1].changes] == ["Status", "Status code"]
    rows = read_records(path)
    assert [(row["Manuscript number"], row["Status"]) for row in rows] == [
        ("SUBTRAK-D-22-00128", "Review Complete"),
        ("OTHER-D-26-00042", "Under Review"),
        ("OTHER-D-26-00042", "Required Reviews Complete"),
    ]
    assert {row["Paper"] for row in rows} == {"Tau"}


def test_changing_the_template_does_not_split_the_history(
    config, fake_client, make_payload, clock
) -> None:
    fake_client.queue(DEMO_UUID, make_payload(), make_payload(Status=4))
    Monitor(config, fake_client, now=clock).run()
    [original] = workbooks(config)

    renamed = replace(
        paper_config(config, "{paper} - {journal_acronym}"),
        submissions=(SubmissionConfig(DEMO_UUID, fields={"paper": "Tau"}),),
    )
    [result] = Monitor(renamed, fake_client, now=clock).run().results

    assert not result.is_new
    assert result.workbook == original
    assert [c.field for c in result.changes] == ["Status", "Status code"]
    assert workbooks(config) == [original]
    assert len(read_records(original)) == 2


def test_disappearing_submission_in_a_shared_workbook(
    config, fake_client, make_payload, clock
) -> None:
    config = paper_config(config)
    fake_client.queue(DEMO_UUID, make_payload(Status=29), SubmissionNotFoundError("HTTP 404"))
    fake_client.queue(OTHER_UUID, other_payload(make_payload))
    monitor = Monitor(config, fake_client, now=clock)
    monitor.run()
    summary = monitor.run()

    gone, other = summary.results
    assert gone.ok
    assert [c.field for c in gone.changes] == ["Status", "Status code"]
    assert other.changes == []
    rows = read_records(workbooks(config)[0])
    assert [(row["Manuscript number"], row["Status"]) for row in rows] == [
        ("SUBTRAK-D-22-00128", "Review Complete"),
        ("OTHER-D-26-00042", "Under Review"),
        ("SUBTRAK-D-22-00128", UNAVAILABLE_STATUS),
    ]
    assert rows[-1]["Paper"] == "Tau"
