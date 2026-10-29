"""Excel storage: workbooks that keep one row per observed state of a submission.

The rows of a submission are identified by its *base* manuscript number (without
the ``R1``, ``R2``... revision suffix) or by its tracking link, so every revision of
a manuscript submitted to one journal is part of the same history.

By default every manuscript gets its own workbook whose name starts with that
number, followed by the journal and the title to make it easy to recognise::

    ABC-D-24-00001 - Journal of Examples - A study of something.xlsx

The name is built from the ``workbook_name`` template (see
:mod:`paper_status_monitor.naming`) and several submissions may share a workbook,
for example all the submissions of one paper to different journals. Workbooks are
found again by their content (see :class:`WorkbookStore`), never by the title, so a
later change of the title or of the template does not split a history.

The history is kept in the ``History`` sheet. With the ``time_in_status`` setting,
a second sheet, ``Time in status``, shows how long each submission spent in each
status (see :func:`write_time_sheet`).
"""

from __future__ import annotations

import logging
import os
import re
import time
from collections.abc import Callable, Iterable, Iterator, Mapping, Sequence
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from paper_status_monitor import timeline
from paper_status_monitor.models import (
    COLUMNS,
    DETECTED_AT_COLUMN,
    MANUSCRIPT_NUMBER_COLUMN,
    NUMERIC_COLUMNS,
    TRACKER_URL_COLUMN,
    base_manuscript_number,
    cell_to_text,
)

log = logging.getLogger(__name__)

#: Name of the worksheet that holds the history.
SHEET_TITLE = "History"

#: Name of the optional worksheet with the time spent in each status
#: (``time_in_status`` setting). It is rewritten after every change.
TIME_SHEET_TITLE = "Time in status"

#: Maximum length of a workbook file name without the extension.
MAX_STEM_LENGTH = 150

_TIME_SHEET_MARK = "Written by paper-status-monitor"
_TIME_SHEET_NOTE = (
    f"{_TIME_SHEET_MARK} after every change of the History sheet: do not edit this sheet, "
    "your changes would be lost. The days of the current status are counted until now."
)
_PERIOD_HEADINGS = ("Manuscript number", "Journal", "Revision", "Status", "From", "To", "Days")
_TOTAL_HEADINGS = ("Manuscript number", "Journal", "Status", "Days", "Periods")
_TIME_SHEET_WIDTHS = (20, 32, 10, 30, 18, 18, 10)
_DATE_TIME_FORMAT = "yyyy-mm-dd hh:mm"

_INVALID_CHARS = re.compile(r'[<>:"/\\|?*\x00-\x1f]')
_HEADER_FONT = Font(bold=True)
_HEADER_FILL = PatternFill("solid", start_color="DDEBF7", end_color="DDEBF7")
_HEADER_ALIGNMENT = Alignment(wrap_text=True, vertical="center")
_COLUMN_WIDTHS = {
    DETECTED_AT_COLUMN: 18,
    "Title": 60,
    "Revision": 10,
    "Status": 26,
    "Status code": 10,
    "Last review activity": 22,
    "Reviews completed": 11,
    "Review invitations accepted": 12,
    "Review invitations sent": 12,
    "Journal": 32,
    "Corresponding author": 24,
    "First author": 24,
    "Date of submission": 22,
    "Manuscript number": 20,
    TRACKER_URL_COLUMN: 45,
}
_REPLACE_ATTEMPTS = 5
_REPLACE_DELAY = 0.5


class WorkbookLockedError(OSError):
    """The workbook could not be written, typically because it is open in Excel."""


def sanitize_filename(text: str) -> str:
    """Remove characters that are not allowed in Windows file names."""
    cleaned = _INVALID_CHARS.sub("", text)
    cleaned = " ".join(cleaned.split())
    return cleaned.strip(" .")


def workbook_name(key: str, journal: str = "", title: str = "") -> str:
    """Return the default file name of a new workbook.

    Args:
        key: Base manuscript number (or the tracking UUID if it is unknown).
        journal: Journal name, added for readability.
        title: Manuscript title, added for readability (truncated if needed).

    Raises:
        ValueError: If *key* is empty after sanitisation.
    """
    clean_key = sanitize_filename(key)
    if not clean_key:
        raise ValueError("A workbook key (manuscript number or UUID) is required.")
    parts = [clean_key, *(p for p in (sanitize_filename(journal), sanitize_filename(title)) if p)]
    stem = " - ".join(parts)
    if len(stem) > MAX_STEM_LENGTH:
        cut = stem[:MAX_STEM_LENGTH]
        if " " in cut[len(clean_key) + 3 :]:
            cut = cut.rsplit(" ", 1)[0]
        stem = cut.rstrip(" .-")
    return f"{stem}.xlsx"


def _workbook_files(directory: Path) -> Iterator[Path]:
    if not directory.is_dir():
        return
    for path in directory.iterdir():
        if (
            path.suffix.lower() == ".xlsx"
            and not path.name.startswith(("~$", "."))
            and path.is_file()
        ):
            yield path


def is_named_after(path: Path, key: str) -> bool:
    """Return ``True`` if the file name of *path* starts with *key*.

    That is the case of ``"<key>.xlsx"``, ``"<key> - <anything>.xlsx"`` and
    sync-client duplicates such as ``"<key> (1).xlsx"`` (the default names).
    """
    prefix = sanitize_filename(key).casefold()
    stem = path.stem.casefold()
    return bool(prefix) and (stem == prefix or stem.startswith((f"{prefix} - ", f"{prefix} (")))


def find_workbook(directory: Path, key: str) -> Path | None:
    """Return the existing workbook named after *key* in *directory*, or ``None``.

    A workbook matches when its name is ``"<key>.xlsx"``, starts with
    ``"<key> - "`` or is a sync-client duplicate such as ``"<key> (1).xlsx"``.
    If several files match, the most recently modified one is returned and a
    warning is logged.
    """
    matches = [path for path in _workbook_files(directory) if is_named_after(path, key)]
    if not matches:
        return None
    return _most_recent(matches, key)


def find_workbook_by_uuid(directory: Path, uuid: str) -> Path | None:
    """Return the workbook whose ``Tracker URL`` column mentions *uuid*, or ``None``."""
    needle = uuid.casefold()
    for path in _workbook_files(directory):
        try:
            records = read_records(path)
        except Exception as exc:  # unreadable or foreign workbook: skip it
            log.debug("Skipping %s while looking for %s: %s", path.name, uuid, exc)
            continue
        if any(needle in record.get(TRACKER_URL_COLUMN, "").casefold() for record in records):
            return path
    return None


def _most_recent(paths: list[Path], key: str) -> Path:
    paths = sorted(paths, key=lambda p: p.stat().st_mtime, reverse=True)
    if len(paths) > 1:
        log.warning(
            "Found %d workbooks for %s; using the most recently modified one (%s). "
            "Please merge or delete the others: %s",
            len(paths),
            key,
            paths[0].name,
            ", ".join(p.name for p in paths[1:]),
        )
    return paths[0]


def has_identifier(record: Mapping[str, str]) -> bool:
    """Return ``True`` if a row has a manuscript number or a tracking link."""
    return bool(record.get(MANUSCRIPT_NUMBER_COLUMN) or record.get(TRACKER_URL_COLUMN))


def belongs_to(record: Mapping[str, str], key: str, uuid: str = "") -> bool:
    """Return ``True`` if a workbook row was recorded for a given submission.

    Args:
        record: Workbook row (column label to text, as returned by :func:`read_records`).
        key: Base manuscript number of the submission (may be empty).
        uuid: Tracking UUID of the submission (may be empty).
    """
    number = record.get(MANUSCRIPT_NUMBER_COLUMN, "")
    if number and key and base_manuscript_number(number).casefold() == key.casefold():
        return True
    url = record.get(TRACKER_URL_COLUMN, "")
    return bool(uuid and url) and uuid.casefold() in url.casefold()


@dataclass(frozen=True)
class WorkbookLocation:
    """Where the history of a submission is (or will be) kept.

    Attributes:
        path: The workbook (it may not exist yet).
        key: Base manuscript number of the submission (or its UUID if unknown).
        last_record: Last row recorded for the submission, ``None`` if there is none.
    """

    path: Path
    key: str
    last_record: dict[str, str] | None = None


class WorkbookStore:
    """Find and read the workbooks of a data folder.

    Workbooks are read at most once while they do not change on disk, so looking
    for several submissions in the same folder stays fast.

    Args:
        directory: The data folder.
    """

    def __init__(self, directory: Path) -> None:
        self.directory = directory
        self._cache: dict[Path, tuple[tuple[int, int], list[dict[str, str]]]] = {}
        self._sheets: dict[Path, tuple[tuple[int, int], list[str]]] = {}

    def workbooks(self) -> list[Path]:
        """Return the workbooks of the folder (temporary and lock files excluded)."""
        return sorted(_workbook_files(self.directory))

    def records(self, path: Path) -> list[dict[str, str]]:
        """Return every data row of a workbook (see :func:`read_records`)."""
        stat = path.stat()
        signature = (stat.st_mtime_ns, stat.st_size)
        cached = self._cache.get(path)
        if cached is not None and cached[0] == signature:
            return cached[1]
        records = read_records(path)
        self._cache[path] = (signature, records)
        return records

    def sheet_names(self, path: Path) -> list[str]:
        """Return the names of the sheets of a workbook (read once while it does not change)."""
        stat = path.stat()
        signature = (stat.st_mtime_ns, stat.st_size)
        cached = self._sheets.get(path)
        if cached is not None and cached[0] == signature:
            return cached[1]
        workbook = load_workbook(path, read_only=True)
        try:
            names = list(workbook.sheetnames)
        finally:
            workbook.close()
        self._sheets[path] = (signature, names)
        return names

    def submission_records(self, path: Path, key: str, uuid: str = "") -> list[dict[str, str]]:
        """Return the rows of *path* that belong to a submission, in file order.

        Rows without any identifier (written by very old versions) are attributed to
        the submission only when the workbook is named after its key.
        Unreadable workbooks have no rows.
        """
        named = bool(key) and is_named_after(path, key)
        return [
            record
            for record in self._safe_records(path)
            if belongs_to(record, key, uuid) or (named and not has_identifier(record))
        ]

    def locate(
        self, key: str, uuid: str = "", target_name: str = "", *, hint: str | None = None
    ) -> WorkbookLocation:
        """Find the workbook that holds the history of a submission.

        The search order is: the workbook given by the naming template
        (*target_name*), the workbook used last time (*hint*), the workbooks named
        after *key*, and finally every other workbook. A history is never split: if
        the rows of the submission are found anywhere, that workbook is used even if
        its name does not follow the template (``organize`` renames it).

        Args:
            key: Base manuscript number (or the UUID if the number is unknown).
            uuid: Tracking UUID.
            target_name: File name given by the naming template.
            hint: File name where the submission was found by the previous run.

        Returns:
            The location; ``last_record`` is ``None`` for a submission that has not
            been recorded yet, in which case ``path`` is the workbook to create (or
            to share).
        """
        target = self.directory / (target_name or workbook_name(key))
        named = find_workbook(self.directory, key)
        candidates: list[Path] = []
        for path in (target, self.directory / hint if hint else None, named):
            if path is not None and path not in candidates and path.is_file():
                candidates.append(path)
        for path in candidates:
            rows = self.submission_records(path, key, uuid)
            if rows:
                return WorkbookLocation(path, key, rows[-1])
        found = {
            path: rows
            for path in self.workbooks()
            if path not in candidates and (rows := self.submission_records(path, key, uuid))
        }
        if found:
            path = _most_recent(list(found), key)
            return WorkbookLocation(path, key, found[path][-1])
        if named is not None and named != target and not self._safe_records(named):
            return WorkbookLocation(named, key)  # an empty workbook created for this key
        return WorkbookLocation(target, key)

    def locate_uuid(self, uuid: str, *, hint: str | None = None) -> WorkbookLocation | None:
        """Find the history of a submission from its UUID alone (no tracker data).

        Returns:
            The location in the workbook whose rows mention *uuid* in the
            ``Tracker URL`` column (the most recently modified one if there are
            several), or ``None`` if there is none.
        """
        paths = self.workbooks()
        preferred = self.directory / hint if hint else None
        if preferred in paths:
            paths.remove(preferred)
            paths.insert(0, preferred)
        found: dict[Path, WorkbookLocation] = {}
        for path in paths:
            mentions = [
                record for record in self._safe_records(path) if belongs_to(record, "", uuid)
            ]
            if not mentions:
                continue
            number = mentions[-1].get(MANUSCRIPT_NUMBER_COLUMN, "")
            key = base_manuscript_number(number) or uuid
            rows = self.submission_records(path, key, uuid)
            found[path] = WorkbookLocation(path, key, rows[-1])
            if path == preferred:
                break
        if not found:
            return None
        return found[_most_recent(list(found), uuid)]

    def append(
        self, path: Path, record: Mapping[str, Any], *, time_in_status: bool = False
    ) -> None:
        """Append a row (see :func:`append_record`) and forget the cached rows."""
        self.forget(path)
        append_record(path, record, time_in_status=time_in_status)
        self.forget(path)

    def add_time_sheet(self, path: Path) -> bool:
        """Add the *Time in status* sheet to an existing workbook that does not have it.

        Returns:
            ``True`` if the sheet was added.
        """
        if TIME_SHEET_TITLE in self.sheet_names(path):
            return False
        self.forget(path)
        added = write_time_sheet(path)
        self.forget(path)
        return added

    def forget(self, path: Path | None = None) -> None:
        """Drop the cached rows of *path* (of every workbook if ``None``)."""
        if path is None:
            self._cache.clear()
            self._sheets.clear()
        else:
            self._cache.pop(path, None)
            self._sheets.pop(path, None)

    def _safe_records(self, path: Path) -> list[dict[str, str]]:
        try:
            return self.records(path)
        except Exception as exc:  # unreadable or foreign workbook: no rows
            log.debug("Could not read %s: %s", path.name, exc)
            return []


def _history_sheet(workbook: Workbook) -> Worksheet:
    """Return the history sheet: ``History``, else the first sheet that is not the
    *Time in status* sheet."""
    if SHEET_TITLE in workbook.sheetnames:
        return workbook[SHEET_TITLE]
    others = [sheet for sheet in workbook.worksheets if sheet.title != TIME_SHEET_TITLE]
    return (others or workbook.worksheets)[0]


def read_records(path: Path) -> list[dict[str, str]]:
    """Read every data row of a workbook as a mapping of column label to string."""
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        rows = _history_sheet(workbook).iter_rows(values_only=True)
        header = next(rows, None)
        if not header:
            return []
        labels = [cell_to_text(value) for value in header]
        records = []
        for row in rows:
            values = [cell_to_text(value) for value in row]
            if not any(values):
                continue
            values += [""] * (len(labels) - len(values))
            records.append({label: values[i] for i, label in enumerate(labels) if label})
        return records
    finally:
        workbook.close()


def last_record(path: Path) -> dict[str, str] | None:
    """Return the last data row of a workbook, or ``None`` if it has no data rows."""
    records = read_records(path)
    return records[-1] if records else None


def append_record(path: Path, record: Mapping[str, Any], *, time_in_status: bool = False) -> None:
    """Append one row to the workbook at *path*, creating the file if needed.

    Values are written under the matching header. Columns missing from an older
    workbook are added at the end of the header, so files written by previous
    versions keep working. The file is replaced atomically.

    Args:
        path: The workbook.
        record: Column label to value.
        time_in_status: Add the *Time in status* sheet (see :func:`write_time_sheet`).
            An existing one is always brought up to date.

    Raises:
        WorkbookLockedError: If the file cannot be replaced (e.g. open in Excel).
    """
    if path.exists():
        workbook = load_workbook(path)
        sheet = _history_sheet(workbook)
    else:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = SHEET_TITLE
    try:
        header = _ensure_header(sheet, [*COLUMNS, *record.keys()])
        row = _last_used_row(sheet) + 1
        _write_row(sheet, row, header, record)
        _finish_sheet(sheet, header, row)
        if time_in_status or TIME_SHEET_TITLE in workbook.sheetnames:
            _write_time_sheet(workbook, sheet, path.name)
        _save_atomic(workbook, path)
    finally:
        workbook.close()


def merge_workbooks(target: Path, sources: Sequence[Path], *, time_in_status: bool = False) -> int:
    """Combine the history of several workbooks into *target*.

    The rows of *target* (if it exists) and of every source are merged in
    chronological order (``Detected at``). The order of the rows of each workbook
    is kept, and a row without a date stays after the rows that preceded it. Every
    column is kept. *target* is used as the base when it exists (its other sheets
    and column widths are preserved); otherwise the first source is. The sources
    are **not** deleted.

    Args:
        target: The merged workbook.
        sources: Workbooks whose rows are added.
        time_in_status: Add the *Time in status* sheet; an existing one is always
            rewritten from the merged rows.

    Returns:
        The number of data rows written.

    Raises:
        WorkbookLockedError: If *target* cannot be written.
    """
    if not sources:
        raise ValueError("At least one source workbook is required.")
    base = target if target.exists() else sources[0]
    tables = [_read_table(path) for path in sources if path != base]
    workbook = load_workbook(base)
    try:
        sheet = _history_sheet(workbook)
        tables.insert(0, _sheet_table(sheet))
        labels = [label for header, _ in tables for label in header]
        rows = _chronological([rows for _, rows in tables])
        if sheet.max_row > 1:
            sheet.delete_rows(2, sheet.max_row - 1)
        header = _ensure_header(sheet, [*COLUMNS, *labels])
        for index, record in enumerate(rows, start=2):
            _write_row(sheet, index, header, record)
        _finish_sheet(sheet, header, len(rows) + 1)
        if time_in_status or TIME_SHEET_TITLE in workbook.sheetnames:
            _write_time_sheet(workbook, sheet, target.name)
        _save_atomic(workbook, target)
    finally:
        workbook.close()
    return len(rows)


def copy_rows(
    source: Path,
    target: Path,
    keep: Callable[[Mapping[str, str]], bool],
    *,
    time_in_status: bool = False,
) -> int:
    """Write the rows of *source* that *keep* accepts to *target*.

    *keep* receives every row as text (like the rows of :func:`read_records`).

    * When *target* is *source*, the other rows are removed from it.
    * When *target* is another existing workbook, the rows are merged into it in
      chronological order, as :func:`merge_workbooks` does.
    * Otherwise *target* is created from a copy of *source*, so that its formatting
      and its other sheets are kept, with only those rows.

    *source* is only modified when it is *target*.

    Returns:
        The number of rows of *source* written to *target*.

    Raises:
        WorkbookLockedError: If *target* cannot be written.
    """
    labels, rows = _read_table(source)
    selected = [
        row for row in rows if keep({label: cell_to_text(value) for label, value in row.items()})
    ]
    merging = target.exists() and target != source and not _same_file(source, target)
    workbook = load_workbook(target if merging else source)
    try:
        sheet = _history_sheet(workbook)
        records = selected
        if merging:
            existing_labels, existing = _sheet_table(sheet)
            records = _chronological([existing, selected])
            labels = [*existing_labels, *labels]
        if sheet.max_row > 1:
            sheet.delete_rows(2, sheet.max_row - 1)
        header = _ensure_header(sheet, [*COLUMNS, *labels])
        for index, record in enumerate(records, start=2):
            _write_row(sheet, index, header, record)
        _finish_sheet(sheet, header, len(records) + 1)
        if time_in_status or TIME_SHEET_TITLE in workbook.sheetnames:
            _write_time_sheet(workbook, sheet, target.name)
        _save_atomic(workbook, target)
    finally:
        workbook.close()
    return len(selected)


def ensure_writable(paths: Iterable[Path]) -> None:
    """Make sure that the existing workbooks among *paths* can be replaced.

    Raises:
        WorkbookLockedError: For the first one that is in use (typically, open in Excel).
    """
    for path in paths:
        if not path.exists():
            continue
        try:
            with path.open("r+b"):
                pass
        except PermissionError as exc:
            raise WorkbookLockedError(
                f"Cannot write {path.name}: the file is in use. Is it open in Excel?"
            ) from exc


def _same_file(first: Path, second: Path) -> bool:
    try:
        return os.path.samefile(first, second)
    except OSError:
        return False


def write_time_sheet(path: Path) -> bool:
    """Write (or rewrite) the *Time in status* sheet of an existing workbook.

    The sheet lists, for every submission of the workbook, the periods during which
    it kept the same status and revision (from, to and days) and the total days in
    each status (see :mod:`~paper_status_monitor.timeline`). The days of the
    current status are an Excel formula (``NOW()``), so they stay up to date
    without rewriting the workbook. The history is not modified.

    Returns:
        ``False`` if the workbook already has a sheet with that name that was not
        written by this program: it is left untouched.

    Raises:
        WorkbookLockedError: If the file cannot be replaced (e.g. open in Excel).
    """
    workbook = load_workbook(path)
    try:
        if not _write_time_sheet(workbook, _history_sheet(workbook), path.name):
            return False
        _save_atomic(workbook, path)
    finally:
        workbook.close()
    return True


def _write_time_sheet(workbook: Workbook, history: Worksheet, name: str = "") -> bool:
    """Build the *Time in status* sheet from the rows of *history*.

    See :func:`write_time_sheet`; *name* identifies the workbook in log messages.
    """
    position = None
    if TIME_SHEET_TITLE in workbook.sheetnames:
        existing = workbook[TIME_SHEET_TITLE]
        if not str(existing.cell(2, 1).value or "").startswith(_TIME_SHEET_MARK):
            log.warning(
                "%s has a sheet named %r that was not written by paper-status-monitor; it "
                "is left as it is. Rename it to get the time spent in each status.",
                name or "A workbook",
                TIME_SHEET_TITLE,
            )
            return False
        position = workbook.sheetnames.index(TIME_SHEET_TITLE)
        workbook.remove(existing)
    sheet = workbook.create_sheet(TIME_SHEET_TITLE, position)
    sheet.cell(1, 1, "Time in each status").font = Font(bold=True, size=14)
    sheet.cell(2, 1, _TIME_SHEET_NOTE).font = Font(italic=True, color="5F6368")
    for column, width in enumerate(_TIME_SHEET_WIDTHS, start=1):
        sheet.column_dimensions[get_column_letter(column)].width = width
    periods = timeline.status_periods(_sheet_table(history)[1])
    if not periods:
        sheet.cell(4, 1, "No dated rows yet.")
        return True
    first = 5
    last = first + len(periods) - 1
    _heading_row(sheet, first - 1, _PERIOD_HEADINGS)
    for row, period in enumerate(periods, start=first):
        sheet.cell(row, 1, period.submission)
        sheet.cell(row, 2, period.journal)
        sheet.cell(row, 3, _to_cell("Revision", period.revision))
        sheet.cell(row, 4, period.status)
        sheet.cell(row, 5, period.start).number_format = _DATE_TIME_FORMAT
        if period.end is not None:
            sheet.cell(row, 6, period.end).number_format = _DATE_TIME_FORMAT
        elif not period.final:
            sheet.cell(row, 6, "now")
        days = period.days()
        if period.ongoing:
            sheet.cell(row, 7, f"=ROUND(NOW()-E{row},1)").number_format = "0.0"
        elif days is not None:
            sheet.cell(row, 7, round(days, 1)).number_format = "0.0"
    title = last + 2
    sheet.cell(title, 1, "Total days in each status").font = Font(bold=True)
    _heading_row(sheet, title + 1, _TOTAL_HEADINGS)

    def span(column: str) -> str:
        return f"${column}${first}:${column}${last}"

    totals = timeline.total_days(periods)
    journals = {period.submission: period.journal for period in periods}
    for row, (submission, status) in enumerate(totals, start=title + 2):
        sheet.cell(row, 1, submission)
        sheet.cell(row, 2, journals.get(submission, ""))
        sheet.cell(row, 3, status)
        criteria = f"{span('A')},A{row},{span('D')},C{row}"
        sheet.cell(row, 4, f"=SUMIFS({span('G')},{criteria})").number_format = "0.0"
        sheet.cell(row, 5, f"=COUNTIFS({criteria})")
    return True


def _heading_row(sheet: Worksheet, row: int, headings: Sequence[str]) -> None:
    for column, heading in enumerate(headings, start=1):
        cell = sheet.cell(row, column, heading)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = _HEADER_ALIGNMENT


def _read_table(path: Path) -> tuple[list[str], list[dict[str, Any]]]:
    workbook = load_workbook(path)
    try:
        return _sheet_table(_history_sheet(workbook))
    finally:
        workbook.close()


def _sheet_table(sheet: Worksheet) -> tuple[list[str], list[dict[str, Any]]]:
    """Return the header and the data rows (raw cell values) of a sheet."""
    rows = sheet.iter_rows(values_only=True)
    first = next(rows, None)
    header = [cell_to_text(value) for value in first] if first else []
    records = []
    for row in rows:
        if all(value in (None, "") for value in row):
            continue
        records.append({label: value for label, value in zip(header, row, strict=False) if label})
    return [label for label in header if label], records


def _chronological(tables: list[list[dict[str, Any]]]) -> list[dict[str, Any]]:
    keyed = []
    for table_index, rows in enumerate(tables):
        latest = datetime.min
        for row_index, row in enumerate(rows):
            moment = timeline.as_datetime(row.get(DETECTED_AT_COLUMN))
            if moment is not None and moment > latest:
                latest = moment
            keyed.append(((latest, table_index, row_index), row))
    keyed.sort(key=lambda item: item[0])
    return [row for _, row in keyed]


def _ensure_header(sheet: Worksheet, labels: list[str]) -> list[str]:
    """Make sure every label is present in the first row; return the header."""
    if sheet.max_row == 1 and sheet.max_column == 1 and sheet.cell(1, 1).value is None:
        header: list[str] = []
    else:
        header = [cell_to_text(cell.value) for cell in sheet[1]]
        while header and not header[-1]:
            header.pop()
    for label in labels:
        if label in header:
            continue
        header.append(label)
        column = len(header)
        cell = sheet.cell(row=1, column=column, value=label)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = _HEADER_ALIGNMENT
        width = _COLUMN_WIDTHS.get(label, 18)
        sheet.column_dimensions[get_column_letter(column)].width = width
    sheet.row_dimensions[1].height = 30
    return header


def _write_row(sheet: Worksheet, row: int, header: list[str], record: Mapping[str, Any]) -> None:
    for label, value in record.items():
        cell = sheet.cell(row=row, column=header.index(label) + 1, value=_to_cell(label, value))
        if isinstance(value, datetime):
            cell.number_format = "yyyy-mm-dd hh:mm"
        elif label == TRACKER_URL_COLUMN and isinstance(value, str) and value:
            cell.hyperlink = value
            cell.style = "Hyperlink"


def _finish_sheet(sheet: Worksheet, header: list[str], last_row: int) -> None:
    last_column = get_column_letter(len(header))
    sheet.auto_filter.ref = f"A1:{last_column}{max(last_row, 1)}"
    sheet.freeze_panes = "A2"


def _last_used_row(sheet: Worksheet) -> int:
    for row in range(sheet.max_row, 0, -1):
        if any(cell.value not in (None, "") for cell in sheet[row]):
            return row
    return 0


def _to_cell(label: str, value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, (datetime, date, int, float)) and not isinstance(value, bool):
        return value
    text = str(value)
    if not text:
        return None
    if label in NUMERIC_COLUMNS and text.isascii() and text.isdigit():
        return int(text)
    return text


def _save_atomic(workbook: Workbook, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_name(f".{path.stem}.tmp-{os.getpid()}.xlsx")
    try:
        workbook.save(temporary)
        for attempt in range(1, _REPLACE_ATTEMPTS + 1):
            try:
                os.replace(temporary, path)
                return
            except PermissionError as exc:
                if attempt == _REPLACE_ATTEMPTS:
                    raise WorkbookLockedError(
                        f"Cannot write {path.name}: the file is in use. Is it open in Excel?"
                    ) from exc
                time.sleep(_REPLACE_DELAY)
    finally:
        if temporary.exists():
            try:
                temporary.unlink()
            except OSError:
                log.debug("Could not remove temporary file %s", temporary)
